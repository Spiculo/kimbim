# SPDX-License-Identifier: GPL-3.0-or-later
"""Schedules and quantities (doc 14 §15, D77): built-in and custom schemes in the project settings,
quantities checked on the sample house's known geometry, merging identical items, editing cells, export."""
import csv
import json
import sys
from pathlib import Path

import ifcopenshell.api.pset
import pytest

from kimcad import app_bridge as b
from kimcad import schedules as S
from kimcad import settings

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "tools"))
import make_house  # noqa: E402


def call(fn, *args):
    return json.loads(getattr(b, fn)(*args))


@pytest.fixture(scope="module")
def house_file(tmp_path_factory):
    path = tmp_path_factory.mktemp("house") / "house.ifc"
    make_house.build(str(path))
    return str(path)


@pytest.fixture
def house(house_file):
    assert call("open_model", house_file)["ok"]


def table(sid, merge=False):
    d = call("schedule", sid, "true" if merge else "false")
    assert d["ok"], d
    keys = [c["key"] for c in d["columns"]]
    rows = [dict(zip(keys, r["values"]), _guids=r["guids"], _count=r["count"], _varies=dict(zip(keys, r["varies"])))
            for r in d["rows"]]
    return d, rows, dict(zip(keys, d["totals"]))


def test_builtin_schemes(house):
    ids = [s["id"] for s in call("schedule_schemes")["schemes"]]
    assert ids == ["doors", "windows", "walls", "slabs", "columns", "beams", "rooms"]


def test_door_and_window_sizes(house):
    _, doors, totals = table("doors")
    assert len(doors) == 2 and [d["info:level"] for d in doors] == ["Level 1", "Level 2"]  # sorted by level
    assert all((d["quantity:width"], d["quantity:height"]) == (0.9, 2.1) for d in doors)
    assert totals["quantity:area"] == pytest.approx(2 * 0.9 * 2.1)
    _, windows, totals = table("windows")
    assert len(windows) == 8 and totals["quantity:area"] == pytest.approx(8 * 1.44)


def test_wall_quantities_on_level_1(house):
    _, walls, _ = table("walls")
    l1 = [w for w in walls if w["info:level"] == "Level 1"]
    # A 10 x 8 m loop (axes) of 300 mm walls, 3 m high; the front wall has the door and a window.
    assert sorted(w["quantity:length"] for w in l1) == pytest.approx([8, 8, 10, 10])
    assert all(w["quantity:height"] == pytest.approx(3.0) and w["quantity:thickness"] == pytest.approx(0.3) for w in l1)
    assert sum(w["quantity:gross_side_area"] for w in l1) == pytest.approx(36 * 3)
    openings = 0.9 * 2.1 + 4 * 1.44
    assert sum(w["quantity:opening_area"] for w in l1) == pytest.approx(openings)
    assert sum(w["quantity:net_side_area"] for w in l1) == pytest.approx(108 - openings)
    assert sum(w["quantity:volume"] for w in l1) == pytest.approx((108 - openings) * 0.3)


def test_walls_cut_by_the_roof_are_measured_from_their_shape(house):
    _, walls, _ = table("walls")
    l2 = [w for w in walls if w["info:level"] == "Level 2"]
    assert len(l2) == 4
    for w in l2:  # attached to the gable roof: taller than 3 m on average, consistent areas
        assert w["quantity:height"] > 3.0
        assert w["quantity:net_side_area"] == pytest.approx(w["quantity:gross_side_area"] - w["quantity:opening_area"])
        assert w["quantity:volume"] == pytest.approx(w["quantity:net_side_area"] * 0.3)


def test_slab_column_beam_and_room_quantities(house):
    _, slabs, totals = table("slabs")
    assert len(slabs) == 3 and all(s["quantity:thickness"] == pytest.approx(0.2) for s in slabs)
    assert all(s["quantity:gross_area"] == pytest.approx(10.3 * 8.3) for s in slabs)  # to the outer faces
    l2 = next(s for s in slabs if s["info:level"] == "Level 2")
    assert l2["quantity:opening_area"] == pytest.approx(5.1 * 1.2)  # the stair hole
    assert l2["quantity:area"] == pytest.approx(10.3 * 8.3 - 5.1 * 1.2)
    assert totals["quantity:volume"] == pytest.approx((3 * 10.3 * 8.3 - 5.1 * 1.2) * 0.2)
    _, (column,), _ = table("columns")
    assert (column["quantity:height"], column["quantity:area"], column["quantity:volume"]) == pytest.approx((3.0, 0.09, 0.27))
    _, (beam,), _ = table("beams")
    assert (beam["quantity:length"], beam["quantity:area"], beam["quantity:volume"]) == pytest.approx((9.7, 0.08, 0.776))
    _, rooms, _ = table("rooms")
    assert [r["attribute:Name"] for r in rooms] == ["101", "201"]
    living = rooms[0]  # inside the wall faces: 9.7 x 7.7 m, 3 m high
    assert living["attribute:LongName"] == "Living room"
    assert living["quantity:area"] == pytest.approx(9.7 * 7.7) and living["quantity:perimeter"] == pytest.approx(34.8)
    assert living["quantity:height"] == pytest.approx(3.0) and living["quantity:volume"] == pytest.approx(9.7 * 7.7 * 3)


def test_merge_identical_items(house):
    d, rows, totals = table("windows", merge=True)
    assert len(rows) == 1 and rows[0]["_count"] == 8 and len(rows[0]["_guids"]) == 8
    assert rows[0]["quantity:area"] == pytest.approx(8 * 1.44) and rows[0]["quantity:width"] == 1.2
    assert rows[0]["info:level"] is None and rows[0]["_varies"]["info:level"]  # two levels merged
    assert d["count"] == 8
    _, walls, totals = table("walls", merge=True)
    assert len(walls) == 1 and walls[0]["quantity:length"] == pytest.approx(72)  # 2 x (10 + 8 + 10 + 8)


def test_quantity_sets_win_over_computed_values(house):
    wall = b._model.by_type("IfcWall")[0]

    def action():
        qto = ifcopenshell.api.pset.add_qto(b._model, product=wall, name="Qto_WallBaseQuantities")
        ifcopenshell.api.pset.edit_qto(b._model, qto=qto, properties={"NetSideArea": 12.5, "Length": 4000.0})  # m², mm
        return 1

    b._run(action)
    _, walls, _ = table("walls")
    row = next(w for w in walls if w["_guids"] == [wall.GlobalId])
    assert row["quantity:net_side_area"] == 12.5 and row["quantity:length"] == pytest.approx(4.0)


def test_custom_scheme_in_the_project_settings(house):
    scheme = {"name": "Openings by type", "classes": ["IfcDoor", "IfcWindow"],
              "fields": [{"source": "info", "name": "class"}, {"source": "info", "name": "type"},
                         {"source": "quantity", "name": "area"}, {"source": "quantity", "name": "count"}],
              "group_by": ["info:class", "info:type"]}
    r = call("save_schedule_scheme", json.dumps(scheme))
    assert r["ok"] and r["id"] == "openings-by-type" and r["dirty"]
    assert settings.current()["schedules"]["openings-by-type"]["classes"] == ["IfcDoor", "IfcWindow"]
    listed = next(s for s in call("schedule_schemes")["schemes"] if s["id"] == "openings-by-type")
    assert not listed["builtin"]
    _, rows, totals = table("openings-by-type", merge=True)
    assert [(r["info:class"], r["quantity:count"]) for r in rows] == [("IfcDoor", 2), ("IfcWindow", 8)]
    assert totals["quantity:count"] == 10
    assert not call("save_schedule_scheme", json.dumps({**scheme, "classes": []}))["ok"]
    assert not call("save_schedule_scheme", json.dumps({**scheme, "fields": [{"source": "quantity", "name": "mass"}]}))["ok"]
    # Editing a built-in stores a copy; deleting it brings the default back.
    walls = next(s for s in call("schedule_schemes")["schemes"] if s["id"] == "walls")
    assert call("save_schedule_scheme", json.dumps({**walls, "fields": walls["fields"][:3]}))["ok"]
    assert next(s for s in call("schedule_schemes")["schemes"] if s["id"] == "walls")["modified"]
    assert call("delete_schedule_scheme", "walls")["ok"] and call("delete_schedule_scheme", "openings-by-type")["ok"]
    assert len(next(s for s in call("schedule_schemes")["schemes"] if s["id"] == "walls")["fields"]) == 13
    assert "openings-by-type" not in settings.current()["schedules"]


def test_schedules_merge_as_a_settings_table():
    base = {"plan_cut_height": 1.2, "schedules": {}}
    ours = {**base, "schedules": {"a": {"name": "A"}}}  # this option added a schedule
    theirs = {**base, "plan_cut_height": 1.5}            # the other changed the cut height
    merged, conflicts = settings.merge(base, ours, theirs)
    assert merged == {"plan_cut_height": 1.5, "schedules": {"a": {"name": "A"}}} and conflicts == []


def test_editing_cells_is_one_undo_step(house):
    _, doors, _ = table("doors")
    guids = [g for d in doors for g in d["_guids"]]
    field = {"source": "attribute", "name": "Tag"}
    r = call("set_schedule_value", json.dumps(guids), json.dumps(field), json.dumps("D1"))
    assert r["ok"] and r["changed"] == 2
    assert [d["attribute:Tag"] for d in table("doors")[1]] == ["D1", "D1"]
    r = call("set_schedule_value", json.dumps(guids[:1]), json.dumps({"source": "property", "pset": "Pset_DoorCommon",
                                                                        "name": "FireRating"}), json.dumps("EI30"))
    assert r["ok"]
    assert table("doors")[1][0]["property:Pset_DoorCommon:FireRating"] == "EI30"
    assert call("undo")["ok"]
    assert table("doors")[1][0]["property:Pset_DoorCommon:FireRating"] is None
    assert call("undo")["ok"]
    assert [d["attribute:Tag"] for d in table("doors")[1]] == [None, None]
    assert not call("set_schedule_value", json.dumps(guids), json.dumps({"source": "quantity", "name": "area"}), "1")["ok"]


def test_export_csv_and_xlsx(house, tmp_path):
    out = tmp_path / "walls.csv"
    r = call("export_schedule", "walls", "true", str(out))
    assert r["ok"] and r["rows"] == 1
    with open(out, newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    assert rows[0][:5] == ["Mark", "Name", "Type", "Level", "Length (m)"] and rows[0][-1] == "Count"
    assert rows[1][4] == "72.00" and rows[1][-1] == "8" and rows[-1][0] == "Total"
    import openpyxl

    out = tmp_path / "doors.xlsx"
    assert call("export_schedule", "doors", "false", str(out))["ok"]
    ws = openpyxl.load_workbook(out).active
    assert ws.title == "Doors" and ws.cell(1, 5).value == "Width (m)"
    assert ws.cell(2, 5).value == pytest.approx(0.9) and ws.cell(4, 1).value == "Total"
    assert ws.cell(4, 7).value == pytest.approx(3.78)


def test_fields_and_classes(house):
    fields = {f["key"] for f in call("schedule_fields", json.dumps(["IfcWall"]))["fields"]}
    assert {"attribute:Name", "info:type", "quantity:net_side_area"} <= fields
    assert any(k.startswith("property:") for k in fields)
    classes = {c["name"]: c["count"] for c in call("schedule_classes")["classes"]}
    assert classes["IfcWall"] == 8 and classes["IfcDoor"] == 2 and "IfcOpeningElement" not in classes


def test_no_model():
    b._model = None
    d = call("schedule", "doors", "false")
    assert d["ok"] and d["rows"] == [] and S.compute("nope")["ok"] is False
